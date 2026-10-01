"""Cross-platform replacements for the Bash helper scripts.

The helpers deliberately call :func:`hcrprobeforge.core.main` for every
channel.  This keeps the scientific implementation, file layout, metadata,
and error behavior in one place while making catalogue runs usable on Windows
as well as macOS and Linux.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import hashlib
import io
import json
import os
import re
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable, Sequence

from . import core, references


CHANNELS = ("B1", "B2", "B3", "B4", "B5")
MANIFEST_CHANNEL_HEADERS = {"channel", "channels", "hcr_channel", "hcr_channels"}
ProgressCallback = Callable[[dict[str, object]], None]
ACCESSION_RE = re.compile(r"^[A-Za-z]{2}_[0-9]+(?:\.[0-9]+)?$")
CONTROLLED_OPTIONS = {
    "--channel",
    "--outdir",
    "--accession",
    "--fasta",
    "--qc-only",
    "--plot-only",
}


def utc_timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _option_present(args: Sequence[str], option: str) -> bool:
    return any(item == option or item.startswith(option + "=") for item in args)


def _option_value(args: Sequence[str], option: str, default: str) -> str:
    for index, item in enumerate(args):
        if item == option and index + 1 < len(args):
            return str(args[index + 1])
        if item.startswith(option + "="):
            return item.split("=", 1)[1]
    return default


def _manifest_resume_payload(
    *,
    extra_args: Sequence[str],
    channels: Sequence[str],
    smart_default: bool,
    species: str,
) -> dict[str, object]:
    """Return the invocation state that determines manifest design output.

    The manifest marker is deliberately based on the forwarded design options
    rather than on output paths.  A changed design/QC/curation option must not
    silently reuse a completed job produced by a different invocation.
    """
    return {
        "hcrprobeforge_version": core.__version__,
        "species": str(species),
        "channels": [str(channel).strip().upper() for channel in channels],
        "smart_default": bool(smart_default),
        "extra_args": [str(item) for item in extra_args],
    }


def _manifest_resume_signature(payload: dict[str, object]) -> str:
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _manifest_done_marker_matches(path: Path, signature: str) -> bool:
    """Return whether a completed marker was written for this invocation.

    Empty markers from earlier releases are intentionally treated as stale and
    are rerun once so they acquire the parameter-aware format.
    """
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, ValueError, TypeError):
        return False
    return isinstance(data, dict) and str(data.get("signature") or "") == signature


def _strip_controlled_options(args: Sequence[str]) -> list[str]:
    """Reject options owned by the batch runner rather than silently drifting."""
    cleaned: list[str] = []
    for item in args:
        option = item.split("=", 1)[0]
        if option in CONTROLLED_OPTIONS:
            raise ValueError(f"{item} is controlled by the batch runner")
        cleaned.append(item)
    return cleaned


def _invoke_core(
    argv: list[str],
    *,
    capture: bool,
    progress_callback: ProgressCallback | None = None,
) -> tuple[int, str, str]:
    import subprocess

    stdout = io.StringIO()
    stderr = io.StringIO()
    try:
        with core.progress_scope(progress_callback):
            if capture:
                with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                    code = core.main(argv)
            else:
                code = core.main(argv)
    except core.RunCancelled:
        raise
    except subprocess.CalledProcessError as exc:
        command = [str(item) for item in exc.cmd] if isinstance(exc.cmd, (list, tuple)) else [str(exc.cmd)]
        print(
            core.design_process_failure_message(int(exc.returncode or 1), command),
            file=stderr if capture else sys.stderr,
        )
        code = int(exc.returncode or 1)
    except core.RetryableNCBIError as exc:
        print(f"ERROR: temporary NCBI failure after retries: {exc}", file=stderr if capture else sys.stderr)
        code = 75
    except (RuntimeError, OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=stderr if capture else sys.stderr)
        code = 1
    except SystemExit as exc:
        code = int(exc.code or 0) if isinstance(exc.code, int) else 1
    return code, stdout.getvalue(), stderr.getvalue()


def _notify(progress_callback: ProgressCallback | None, **update: object) -> None:
    """Send best-effort progress updates without affecting scientific runs."""
    if progress_callback is None:
        return
    try:
        progress_callback(update)
    except core.RunCancelled:
        raise
    except Exception:
        pass


def _resolution_accession(channel_root: Path) -> str:
    candidates = sorted(channel_root.rglob("*_transcript_resolution.json"), key=lambda p: p.stat().st_mtime_ns)
    for path in reversed(candidates):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        accession = str(data.get("selected_accession") or "").strip()
        if accession:
            return accession
    return ""


def _prune_empty_directories(root: Path) -> None:
    """Remove empty directories created by a failed current batch job.

    Transcript validation can fail before the core pipeline writes a target
    file. Keeping those empty target/channel folders makes a manifest look as
    if work was produced, so remove only empty descendants and the supplied
    job label directory. The shared workflow/results parent is never passed as
    root.
    """
    root = Path(root)
    if not root.is_dir():
        return
    directories = sorted(
        (path for path in root.rglob("*") if path.is_dir()),
        key=lambda path: len(path.parts),
        reverse=True,
    )
    for path in [*directories, root]:
        try:
            if path.is_dir() and not any(path.iterdir()):
                path.rmdir()
        except OSError:
            # A concurrently written or user-managed directory is left alone.
            continue


def _run_all_channels_impl(
    *,
    gene: str | None = None,
    accession: str | None = None,
    fasta: Path | None = None,
    output_root: Path = Path("hcr_results"),
    extra_args: Sequence[str] = (),
    channels: Sequence[str] = CHANNELS,
    gene_dir: str | None = None,
    capture: bool = False,
    smart_default: bool = True,
    channel_attempts: int | None = None,
    progress_callback: ProgressCallback | None = None,
    output_scoped: bool = False,
    workflow: str = "design",
) -> dict[str, object]:
    """Run the same single-target CLI once for each requested channel."""
    selected_channels = tuple(channels)
    if not selected_channels or any(channel not in CHANNELS for channel in selected_channels):
        raise ValueError(f"channels must be drawn from {', '.join(CHANNELS)}")
    if fasta and (gene or accession):
        raise ValueError("fasta cannot be combined with gene or accession")
    if not fasta and not (gene or accession):
        raise ValueError("provide a gene, accession, or fasta input")

    forwarded = _strip_controlled_options(list(extra_args))
    species = _option_value(forwarded, "--species", "xtr")
    preset = references.get_species_preset(species)
    if preset is not None:
        species = preset.key
    all_transcripts = _option_present(forwarded, "--all-transcripts")
    if all_transcripts and accession:
        raise ValueError("--all-transcripts cannot be combined with an explicit accession")
    if smart_default and not _option_present(forwarded, "--auto-curate-if-needed"):
        forwarded += ["--auto-curate-if-needed"]
    if smart_default and not _option_present(forwarded, "--target-probes"):
        forwarded += ["--target-probes", "20"]

    if channel_attempts is None:
        raw_attempts = os.environ.get("HCRPROBEFORGE_CHANNEL_ATTEMPTS", "3")
        try:
            channel_attempts = int(raw_attempts)
        except ValueError as exc:
            raise ValueError("HCRPROBEFORGE_CHANNEL_ATTEMPTS must be a positive integer") from exc
    if channel_attempts < 1:
        raise ValueError("channel_attempts must be positive")

    label_dir = gene_dir or core.safe_name(gene or accession or (fasta.stem if fasta else "custom_target")) or "custom_target"
    output_root = (
        Path(output_root).expanduser().resolve()
        if output_scoped
        else references.species_run_root(Path(output_root), species, workflow)
    )
    # Treat one all-channel target as a transaction. If a later channel fails,
    # do not leave successful earlier-channel results beside the failed target
    # and make the user mistake a partial run for a complete one.
    label_root = output_root / label_dir
    label_snapshot = core._snapshot_filesystem(label_root)
    output_root.mkdir(parents=True, exist_ok=True)
    resolved_accession = accession or ""
    records: list[dict[str, object]] = []
    all_stdout: list[str] = []
    all_stderr: list[str] = []

    total_channels = len(selected_channels)
    for channel_index, channel in enumerate(selected_channels, start=1):
        core.raise_if_cancelled()
        _notify(
            progress_callback,
            phase="design",
            completed=channel_index - 1,
            total=total_channels,
            message=f"Preparing channel {channel} ({channel_index}/{total_channels})",
            channel=channel,
            channel_index=channel_index,
            channel_total=total_channels,
            channel_fraction=0.0,
        )
        channel_root = output_root / label_dir / f"{label_dir}_{channel}"
        argv: list[str] = []
        if gene:
            argv.append(gene)
        if fasta:
            argv += ["--fasta", str(Path(fasta).expanduser().resolve())]
        if resolved_accession:
            argv += ["--accession", resolved_accession]
        argv += ["--channel", channel, "--outdir", str(channel_root), "--_species-scoped-outdir"]
        argv += forwarded

        def channel_progress(update: dict[str, object]) -> None:
            enriched = dict(update)
            enriched.update(
                {
                    "channel": channel,
                    "channel_index": channel_index,
                    "channel_total": total_channels,
                }
            )
            _notify(progress_callback, **enriched)

        code = 1
        stdout = ""
        stderr = ""
        for attempt in range(1, channel_attempts + 1):
            core.raise_if_cancelled()
            code, stdout, stderr = _invoke_core(
                argv,
                capture=capture,
                progress_callback=channel_progress,
            )
            all_stdout.append(stdout)
            all_stderr.append(stderr)
            if code != 75 or attempt == channel_attempts:
                break
            if capture:
                all_stderr.append(
                    f"WARNING: temporary NCBI failure in {label_dir} {channel}; retrying "
                    f"(attempt {attempt + 1}/{channel_attempts}).\n"
                )
            else:
                print(
                    f"WARNING: temporary NCBI failure in {label_dir} {channel}; retrying "
                    f"(attempt {attempt + 1}/{channel_attempts}).",
                    file=sys.stderr,
                )

        records.append({"channel": channel, "return_code": code, "output_dir": str(channel_root), "stdout": stdout, "stderr": stderr})
        _notify(
            progress_callback,
            phase="design",
            completed=channel_index,
            total=total_channels,
            message=(
                f"Finished channel {channel} ({channel_index}/{total_channels})"
                if code == 0
                else f"Channel {channel} stopped with return code {code}"
            ),
            channel=channel,
            channel_index=channel_index,
            channel_total=total_channels,
            channel_fraction=1.0,
        )
        if code != 0:
            core._remove_new_filesystem_entries(label_root, label_snapshot)
            _prune_empty_directories(label_root)
            return {
                "return_code": code,
                "success": False,
                "gene_dir": label_dir,
                "resolved_accession": resolved_accession,
                "channels": records,
                "stdout": "".join(all_stdout),
                "stderr": "".join(all_stderr),
            }
        if not resolved_accession and gene and not all_transcripts:
            resolved_accession = _resolution_accession(channel_root)
            if not resolved_accession:
                core._remove_new_filesystem_entries(label_root, label_snapshot)
                _prune_empty_directories(label_root)
                return {
                    "return_code": 1,
                    "success": False,
                    "gene_dir": label_dir,
                    "resolved_accession": "",
                    "channels": records,
                    "stdout": "".join(all_stdout),
                    "stderr": "".join(all_stderr) + "ERROR: completed first channel did not record a selected accession.\n",
                }

    return {
        "return_code": 0,
        "success": True,
        "gene_dir": label_dir,
        "resolved_accession": resolved_accession,
        "channels": records,
        "stdout": "".join(all_stdout),
        "stderr": "".join(all_stderr),
    }


def run_all_channels(
    *,
    gene: str | None = None,
    accession: str | None = None,
    fasta: Path | None = None,
    output_root: Path = Path("hcr_results"),
    extra_args: Sequence[str] = (),
    channels: Sequence[str] = CHANNELS,
    gene_dir: str | None = None,
    capture: bool = False,
    smart_default: bool = True,
    channel_attempts: int | None = None,
    progress_callback: ProgressCallback | None = None,
    output_scoped: bool = False,
    workflow: str = "design",
    cache_root: Path | None = None,
) -> dict[str, object]:
    """Run one target through the same core CLI for each requested channel.

    The worker temporarily points the core cache at the project-level cache
    tree.  This keeps cache placement deterministic without leaking an
    environment override into a later webapp run in the same process.
    """
    forwarded = list(extra_args)
    species = _option_value(forwarded, "--species", "xtr")
    preset = references.get_species_preset(species)
    if preset is not None:
        species = preset.key
    project_root = Path(output_root).expanduser().resolve()
    cache_path = Path(cache_root).expanduser().resolve() if cache_root else references.species_cache_root(project_root, species, workflow)
    original_cache = os.environ.get("HCRPROBEFORGE_CACHE_DIR")
    original_species = os.environ.get("HCRPROBEFORGE_SPECIES")
    os.environ["HCRPROBEFORGE_CACHE_DIR"] = str(cache_path)
    os.environ["HCRPROBEFORGE_SPECIES"] = species
    try:
        return _run_all_channels_impl(
            gene=gene,
            accession=accession,
            fasta=fasta,
            output_root=output_root,
            extra_args=forwarded,
            channels=channels,
            gene_dir=gene_dir,
            capture=capture,
            smart_default=smart_default,
            channel_attempts=channel_attempts,
            progress_callback=progress_callback,
            output_scoped=output_scoped,
            workflow=workflow,
        )
    finally:
        if original_cache is None:
            os.environ.pop("HCRPROBEFORGE_CACHE_DIR", None)
        else:
            os.environ["HCRPROBEFORGE_CACHE_DIR"] = original_cache
        if original_species is None:
            os.environ.pop("HCRPROBEFORGE_SPECIES", None)
        else:
            os.environ["HCRPROBEFORGE_SPECIES"] = original_species


def _split_manifest_line(line: str) -> list[str]:
    if "\t" in line:
        return [item.strip() for item in line.split("\t")]
    if "," in line:
        return [item.strip() for item in next(csv.reader([line]))]
    # Keep simple whitespace-separated manifests convenient while allowing
    # the headerless three-column form: gene, accession, channel.
    return line.split()


def _is_header(values: Sequence[str]) -> bool:
    lowered = {value.strip().lower() for value in values}
    return bool(
        lowered
        & (
            {
                "gene",
                "gene_symbol",
                "gene name",
                "gene_name",
                "transcript",
                "transcript_accession",
                "accession",
                "refseq_accession",
            }
            | MANIFEST_CHANNEL_HEADERS
        )
    )


def _normalise_manifest_header(values: Sequence[str]) -> list[str]:
    normalized: list[str] = []
    for item in values:
        key = item.strip().lower().replace("-", "_").replace(" ", "_")
        if key in MANIFEST_CHANNEL_HEADERS:
            key = "channel"
        normalized.append(key)
    return normalized


def _manifest_row_from_values(
    values: Sequence[str], line_number: int, header: list[str] | None
) -> dict[str, object]:
    if header is not None:
        mapping = {
            key: values[index].strip() if index < len(values) else ""
            for index, key in enumerate(header)
        }
        gene = mapping.get("gene_symbol") or mapping.get("gene") or mapping.get("gene_name") or ""
        accession = (
            mapping.get("accession")
            or mapping.get("transcript")
            or mapping.get("transcript_accession")
            or mapping.get("refseq_accession")
            or ""
        )
        channel = mapping.get("channel") or ""
    else:
        gene = values[0].strip() if values else ""
        accession = ""
        channel = ""
        if len(values) > 1:
            second = values[1].strip()
            if second.upper() in CHANNELS:
                channel = second
            elif re.fullmatch(r"B[0-9]+", second.upper()):
                # Treat an out-of-range B-number as a channel so a typo such
                # as ``B9`` gets a useful channel error instead of the less
                # helpful "invalid accession" message.
                channel = second
            else:
                accession = second
        if len(values) > 2:
            # Headerless three-column files use gene, accession, channel.
            # Extra columns are deliberately ignored rather than becoming
            # accidental design inputs.
            accession = values[1].strip()
            channel = values[2].strip()
    return {
        "line_number": line_number,
        "gene_symbol": gene,
        "accession": accession,
        "channel": channel,
    }


def _manifest_spreadsheet_rows(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on installation
        raise ValueError(
            "Reading an Excel manifest requires openpyxl. Install it with: "
            "python -m pip install openpyxl."
        ) from exc

    rows: list[dict[str, object]] = []
    header: list[str] | None = None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for line_number, raw_values in enumerate(workbook.active.iter_rows(values_only=True), 1):
            values = ["" if value is None else str(value).strip() for value in raw_values]
            if not any(values):
                continue
            if header is None and _is_header(values):
                header = _normalise_manifest_header(values)
                continue
            rows.append(_manifest_row_from_values(values, line_number, header))
    finally:
        workbook.close()
    return rows


def _manifest_rows(path: Path) -> list[dict[str, object]]:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _manifest_spreadsheet_rows(path)
    if path.suffix.lower() == ".xls":
        raise ValueError(".xls manifests are not supported; save the workbook as .xlsx or use CSV/TSV.")

    rows: list[dict[str, object]] = []
    header: list[str] | None = None
    for line_number, raw in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        stripped = raw.strip()
        if not stripped or stripped.startswith("#"):
            continue
        values = _split_manifest_line(raw.rstrip("\r"))
        if header is None and _is_header(values):
            header = _normalise_manifest_header(values)
            continue
        rows.append(_manifest_row_from_values(values, line_number, header))
    return rows


def _preflight_all_rows(path: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    raw_rows = _manifest_rows(Path(path))
    seen_jobs: set[tuple[str, str, str]] = set()
    accession_labels: defaultdict[str, set[str]] = defaultdict(set)
    gene_dirs: defaultdict[str, set[str]] = defaultdict(set)
    rows: list[dict[str, object]] = []
    for row in raw_rows:
        gene = str(row["gene_symbol"]).strip()
        accession = str(row["accession"]).strip()
        channel = str(row.get("channel") or "").strip().upper()
        gene_dir = core.safe_name(gene)
        status = "valid"
        reasons: list[str] = []
        if not gene:
            status = "invalid"
            reasons.append("blank gene symbol")
        if accession and not ACCESSION_RE.fullmatch(accession):
            status = "invalid"
            reasons.append("invalid accession format")
        if channel and channel not in CHANNELS:
            status = "invalid"
            reasons.append(f"invalid channel {channel!r}; expected B1, B2, B3, B4, or B5")
        key = (gene, accession, channel)
        if status == "valid" and key in seen_jobs:
            status = "duplicate"
            reasons.append("exact duplicate gene+accession+channel row")
        if status == "valid":
            seen_jobs.add(key)
            gene_dirs[gene_dir].add(gene)
            if accession:
                accession_labels[accession].add(gene)
        if gene and gene_dir != gene:
            reasons.append(f"normalized gene directory: {gene} -> {gene_dir}")
        row.update({"channel": channel, "gene_dir": gene_dir, "status": status, "reason": "; ".join(reasons)})
        rows.append(row)

    collisions: list[dict[str, object]] = []
    collision_dirs = {directory: labels for directory, labels in gene_dirs.items() if len(labels) > 1}
    for row in rows:
        if row["status"] == "valid" and str(row["gene_dir"]) in collision_dirs:
            labels = sorted(collision_dirs[str(row["gene_dir"])])
            row["status"] = "directory_collision"
            row["reason"] = f"different gene symbols normalize to {row['gene_dir']}: {', '.join(labels)}"
            collisions.append(row)

    for row in rows:
        accession = str(row["accession"])
        if row["status"] == "valid" and accession and len(accession_labels[accession]) > 1:
            row["status"] = "shared_accession"
            row["reason"] = "same accession supplied under multiple gene labels"

    return rows, collisions


def preflight_manifest(path: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """Validate a flexible gene/accession list before any design starts."""
    rows, collisions = _preflight_all_rows(Path(path))
    jobs = [row for row in rows if row["status"] in {"valid", "shared_accession"}]
    return jobs, collisions


def manifest_channel_collisions(
    rows: Sequence[dict[str, object]], channels: Sequence[str]
) -> list[dict[str, object]]:
    """Find rows that would write the same target/channel output tree.

    A blank row channel expands to the command-level channel selection.  An
    explicit channel must not overlap that expansion for the same gene and
    accession (or a gene-only request), otherwise a later row could overwrite
    an earlier result while still looking like an independent job.
    """
    selected = tuple(str(channel).strip().upper() for channel in channels)
    seen: dict[tuple[str, str, str], dict[str, object]] = {}
    collisions: list[dict[str, object]] = []
    for row in rows:
        if row.get("status") not in {"valid", "shared_accession"}:
            continue
        gene_dir = str(row.get("gene_dir") or "")
        accession = str(row.get("accession") or "")
        explicit = str(row.get("channel") or "").strip().upper()
        row_channels = (explicit,) if explicit else selected
        for channel in row_channels:
            key = (gene_dir, accession, channel)
            prior = seen.get(key)
            if prior is not None:
                collisions.append(
                    {
                        **row,
                        "channel": channel,
                        "reason": (
                            f"same gene/accession/channel output as line {prior.get('line_number')} "
                            f"({prior.get('gene_symbol')})"
                        ),
                    }
                )
            else:
                seen[key] = row
    return collisions


PREFLIGHT_FIELDS = ["line_number", "gene_symbol", "gene_dir", "accession", "channel", "status", "reason"]


def write_preflight(path: Path, rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=PREFLIGHT_FIELDS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in PREFLIGHT_FIELDS} for row in rows)


def _safe_job_name(gene_dir: str, accession: str, channel: str | None = None) -> str:
    suffix = core.safe_name(accession) if accession else "gene_only"
    # Preserve legacy marker names for manifests without a channel column.
    # Explicit per-row channels get their own marker/log namespace so the same
    # accession can safely appear once for B1 and once for B2.
    channel_suffix = f"__{core.safe_name(channel)}" if channel else ""
    return f"{gene_dir}__{suffix}{channel_suffix}"


def _as_int(value: object, default: int = 0) -> int:
    try:
        return int(value) if value not in (None, "") else default
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default


def _as_float(value: object) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _bool_text(value: object) -> str:
    return "yes" if bool(value) else "no"


CHANNEL_OUTCOME_FIELDS = [
    "timestamp", "gene_symbol", "accession", "channel", "pipeline_status", "report_status", "design_outcome",
    "sufficient_vs_capacity", "final_probe_pairs", "requested_target", "canonical_52nt_capacity", "deep_plan_capacity",
    "capacity_aware_goal", "fraction_of_capacity_goal", "capacity_limited_transcript", "annotation_status",
    "annotation_cache_status", "transcript_cache_status", "transcript_type", "five_prime_utr_start", "five_prime_utr_end",
    "cds_start", "cds_end", "three_prime_utr_start", "three_prime_utr_end", "selected_5utr_probes", "selected_cds_probes",
    "selected_3utr_probes", "selected_boundary_probes", "selected_noncoding_probes", "selected_unannotated_probes",
    "qc_stringency", "adaptive_fallback_triggered", "final_dimer_review_flags", "hairpin_review_count",
    "hairpin_strong_review_count", "self_dimer_review_count", "self_dimer_strong_review_count", "same_pair_heterodimer_flags",
    "cross_pair_heterodimer_flags", "coverage_enabled", "occupied_bins", "probe_span_fraction", "maximum_untargeted_gap_nt",
    "report_json", "log",
]
SUMMARY_FIELDS = [
    "timestamp", "gene_symbol", "accession", "channels", "pipeline_status", "overall_design_outcome", "channels_with_reports",
    "channels_requested_target_met", "channels_capacity_sufficient", "channels_partial_below_capacity_goal",
    "channels_with_no_acceptable_probes", "channels_with_no_geometric_capacity", "channels_missing_reports",
    "channels_with_transcript_annotation", "channels_transcript_cache_hits", "annotation_status", "cds_start", "cds_end",
    "minimum_final_probe_pairs", "maximum_final_probe_pairs", "log",
]


def _read_current_rows(path: Path, fields: list[str]) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if not reader.fieldnames:
            return []
        # Read older manifests after a schema extension and fill new columns
        # with blanks.  This keeps existing status/outcome rows resumable.
        return [{field: str(row.get(field) or "") for field in fields} for row in reader]


def _write_current_row_updates(
    path: Path,
    fields: list[str],
    existing: list[dict[str, str]],
    updates: dict[tuple[str, ...], list[dict[str, object]]],
) -> None:
    """Write current outcome rows once per manifest run.

    The former per-job read/ rewrite pattern made a large manifest quadratic in
    the number of genes. Keeping updates in memory preserves the same current
    row semantics while making runtime proportional to the manifest size.
    """
    # A row can now be represented once per channel.  Remove all old outcome
    # rows for the same gene/accession before writing the current rows so a
    # changed channel assignment cannot leave stale channel records behind.
    replacement_keys = {(key[0], key[1]) for key in updates}
    rows = [
        row
        for row in existing
        if (row.get("gene_symbol", ""), row.get("accession", "")) not in replacement_keys
    ]
    for new_rows in updates.values():
        rows.extend({field: str(row.get(field, "")) for field in fields} for row in new_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


STATUS_FIELDS = ["timestamp", "gene_symbol", "accession", "channel", "status", "log"]


def _ensure_status_table(path: Path) -> None:
    """Create or upgrade the manifest status table without losing history."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_text("\t".join(STATUS_FIELDS) + "\n", encoding="utf-8")
        return
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return
    if not lines:
        path.write_text("\t".join(STATUS_FIELDS) + "\n", encoding="utf-8")
        return
    header = lines[0].split("\t")
    if header == STATUS_FIELDS:
        return
    old_fields = ["timestamp", "gene_symbol", "accession", "status", "log"]
    if header != old_fields:
        return
    upgraded = ["\t".join(STATUS_FIELDS)]
    for line in lines[1:]:
        if not line.strip():
            continue
        values = line.split("\t")
        if len(values) >= len(old_fields):
            values = values[:3] + [""] + values[3:5]
        upgraded.append("\t".join(values))
    path.write_text("\n".join(upgraded) + "\n", encoding="utf-8")


def _locate_report(project_root: Path, species: str, gene_dir: str, accession: str, channel: str) -> tuple[Path | None, dict[str, object] | None]:
    # New manifest runs pass their explicit results directory.  The fallback
    # keeps resumability for the older catalogue layout without making that
    # layout the destination for new runs.
    base = Path(project_root)
    if (base / "results").is_dir():
        base = base / "results"
    if (base / "hcr_results").is_dir():
        base = base / "hcr_results" / references.species_component(species)
    root = base / gene_dir / f"{gene_dir}_{channel}"
    matches: list[tuple[Path, dict[str, object]]] = []
    if not root.exists():
        return None, None
    for path in root.rglob("*_auto_curate_report.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        recorded = str(data.get("accession") or "")
        if not accession or recorded == accession:
            matches.append((path, data))
    if not matches:
        return None, None
    return max(matches, key=lambda item: item[0].stat().st_mtime_ns)


def collect_design_outcomes(
    *,
    project_root: Path,
    species: str,
    gene_symbol: str,
    gene_dir: str,
    accession: str,
    pipeline_status: str,
    log_file: Path,
    channels: Sequence[str] = CHANNELS,
) -> tuple[list[dict[str, object]], dict[str, object]]:
    """Rebuild current channel and transcript summaries from JSON reports."""
    timestamp = utc_timestamp()
    rows: list[dict[str, object]] = []
    for channel in channels:
        report_path, data = _locate_report(project_root, species, gene_dir, accession, channel)
        if data is None:
            rows.append({
                "timestamp": timestamp, "gene_symbol": gene_symbol, "accession": accession, "channel": channel,
                "pipeline_status": pipeline_status, "report_status": "missing", "design_outcome": "missing_report",
                "sufficient_vs_capacity": "unknown", "report_json": "", "log": str(log_file),
            })
            continue
        final_pairs = _as_int(data.get("final_probe_pairs"))
        requested = _as_int(data.get("target_probes"))
        capacity = data.get("capacity") or {}
        canonical = _as_int(capacity.get("canonical_52nt_capacity"))
        deep = _as_int(capacity.get("deep_plan_capacity"))
        goal = _as_int(capacity.get("adaptive_goal"), min(requested, deep) if requested else deep)
        limited = bool(requested and requested > deep)
        if deep <= 0:
            outcome, sufficient = "no_geometric_capacity", "not_applicable"
        elif requested and final_pairs >= requested:
            outcome, sufficient = "requested_target_met", "yes"
        elif goal > 0 and final_pairs >= goal:
            outcome, sufficient = ("capacity_limited_goal_met" if limited else "capacity_goal_met"), "yes"
        elif final_pairs <= 0:
            outcome, sufficient = "no_acceptable_probes", "no"
        else:
            outcome, sufficient = "partial_below_capacity_goal", "no"
        final_qc = data.get("final_oligo_qc") or {}
        coverage = data.get("coverage") or {}
        final_metrics = coverage.get("final_metrics") or {}
        annotation = data.get("transcript_annotation") or {}
        resolution = data.get("transcript_resolution") or {}
        transcript_cache = resolution.get("transcript_cache") or {}
        five = annotation.get("five_prime_utr") or {}
        cds = annotation.get("cds") or {}
        three = annotation.get("three_prime_utr") or {}
        counts = data.get("selected_probe_region_counts") or {}
        boundary = sum(_as_int(v) for k, v in counts.items() if "boundary" in str(k))
        unannotated = sum(_as_int(v) for k, v in counts.items() if str(k) in {"unannotated", "partial_or_unannotated"} or "partial" in str(k))
        fraction = f"{final_pairs / goal:.4f}" if goal > 0 else ""
        rows.append({
            "timestamp": timestamp, "gene_symbol": gene_symbol, "accession": accession, "channel": channel,
            "pipeline_status": pipeline_status, "report_status": str(data.get("status", "present")), "design_outcome": outcome,
            "sufficient_vs_capacity": sufficient, "final_probe_pairs": final_pairs, "requested_target": requested,
            "canonical_52nt_capacity": canonical, "deep_plan_capacity": deep, "capacity_aware_goal": goal,
            "fraction_of_capacity_goal": fraction, "capacity_limited_transcript": _bool_text(limited),
            "annotation_status": str(annotation.get("status", "")), "annotation_cache_status": str(annotation.get("cache_status", "")),
            "transcript_cache_status": str(transcript_cache.get("status", "")), "transcript_type": str(annotation.get("transcript_type", "")),
            "five_prime_utr_start": five.get("start", ""), "five_prime_utr_end": five.get("end", ""),
            "cds_start": cds.get("start", ""), "cds_end": cds.get("end", ""), "three_prime_utr_start": three.get("start", ""),
            "three_prime_utr_end": three.get("end", ""), "selected_5utr_probes": _as_int(counts.get("5UTR")),
            "selected_cds_probes": _as_int(counts.get("CDS")), "selected_3utr_probes": _as_int(counts.get("3UTR")),
            "selected_boundary_probes": boundary, "selected_noncoding_probes": _as_int(counts.get("noncoding")),
            "selected_unannotated_probes": unannotated, "qc_stringency": str(data.get("qc_stringency", "")),
            "adaptive_fallback_triggered": _bool_text((data.get("adaptive_curation") or {}).get("triggered", False)),
            "final_dimer_review_flags": _as_int(data.get("final_dimer_review_flags")),
            "hairpin_review_count": _as_int(final_qc.get("hairpin_REVIEW")), "hairpin_strong_review_count": _as_int(final_qc.get("hairpin_STRONG_REVIEW")),
            "self_dimer_review_count": _as_int(final_qc.get("self_dimer_REVIEW")), "self_dimer_strong_review_count": _as_int(final_qc.get("self_dimer_STRONG_REVIEW")),
            "same_pair_heterodimer_flags": _as_int(final_qc.get("same_pair_P1_P2_heterodimers_flagged")),
            "cross_pair_heterodimer_flags": _as_int(final_qc.get("cross_pair_heterodimers_flagged")),
            "coverage_enabled": _bool_text(coverage.get("enabled_for_selection", False)),
            "occupied_bins": _as_int(final_metrics.get("occupied_bins")) if final_metrics else "",
            "probe_span_fraction": f"{_as_float(final_metrics.get('probe_span_fraction')):.6f}" if _as_float(final_metrics.get("probe_span_fraction")) is not None else "",
            "maximum_untargeted_gap_nt": _as_int(final_metrics.get("maximum_untargeted_gap_nt")) if final_metrics else "",
            "report_json": str(report_path), "log": str(log_file),
        })

    reports = [row for row in rows if row.get("report_status") != "missing"]
    outcomes = [str(row.get("design_outcome", "")) for row in rows]
    final_counts = [_as_int(row.get("final_probe_pairs")) for row in reports]
    capacity_sufficient = sum(row.get("sufficient_vs_capacity") == "yes" for row in rows)
    annotated = sum(row.get("annotation_status") in {"resolved", "resolved_partial", "noncoding_no_cds"} for row in rows)
    cache_hits = sum(row.get("transcript_cache_status") == "hit" for row in rows)
    cds_starts = {row.get("cds_start") for row in reports if row.get("cds_start") not in (None, "")}
    cds_ends = {row.get("cds_end") for row in reports if row.get("cds_end") not in (None, "")}
    if outcomes.count("requested_target_met") == len(channels):
        overall = "all_channels_requested_target_met"
    elif capacity_sufficient == len(channels):
        overall = "all_channels_capacity_sufficient"
    elif not reports:
        overall = "no_channel_reports"
    elif outcomes.count("missing_report"):
        overall = "incomplete_channel_reports"
    elif outcomes.count("partial_below_capacity_goal") or outcomes.count("no_acceptable_probes"):
        overall = "one_or_more_channels_below_capacity_goal"
    elif outcomes.count("no_geometric_capacity"):
        overall = "one_or_more_channels_have_no_geometric_capacity"
    else:
        overall = "mixed_channel_outcomes"
    summary = {
        "timestamp": timestamp, "gene_symbol": gene_symbol, "accession": accession, "channels": ",".join(channels), "pipeline_status": pipeline_status,
        "overall_design_outcome": overall, "channels_with_reports": len(reports),
        "channels_requested_target_met": outcomes.count("requested_target_met"), "channels_capacity_sufficient": capacity_sufficient,
        "channels_partial_below_capacity_goal": outcomes.count("partial_below_capacity_goal"),
        "channels_with_no_acceptable_probes": outcomes.count("no_acceptable_probes"), "channels_with_no_geometric_capacity": outcomes.count("no_geometric_capacity"),
        "channels_missing_reports": outcomes.count("missing_report"), "channels_with_transcript_annotation": annotated,
        "channels_transcript_cache_hits": cache_hits, "annotation_status": ",".join(sorted({str(row.get("annotation_status")) for row in reports if row.get("annotation_status")})),
        "cds_start": next(iter(cds_starts)) if len(cds_starts) == 1 else "", "cds_end": next(iter(cds_ends)) if len(cds_ends) == 1 else "",
        "minimum_final_probe_pairs": min(final_counts) if final_counts else "", "maximum_final_probe_pairs": max(final_counts) if final_counts else "", "log": str(log_file),
    }
    return rows, summary


def _collect_rejected_outcomes(
    *,
    gene_symbol: str,
    accession: str,
    pipeline_status: str,
    report_status: str,
    log_file: Path,
    channels: Sequence[str],
) -> tuple[list[dict[str, object]], dict[str, object]]:
    """Create bookkeeping rows for a row rejected before scientific output.

    A row rejected before scientific output must not be allowed to rediscover
    an older report and present it as the result of the current invocation.
    These rows are intentionally report-free and contain only manifest status.
    """
    timestamp = utc_timestamp()
    rows = [
        {
            "timestamp": timestamp,
            "gene_symbol": gene_symbol,
            "accession": accession,
            "channel": channel,
            "pipeline_status": pipeline_status,
            "report_status": report_status,
            "design_outcome": pipeline_status,
            "sufficient_vs_capacity": "not_applicable",
            "report_json": "",
            "log": str(log_file),
        }
        for channel in channels
    ]
    summary = {
        "timestamp": timestamp,
        "gene_symbol": gene_symbol,
        "accession": accession,
        "channels": ",".join(channels),
        "pipeline_status": pipeline_status,
        "overall_design_outcome": pipeline_status,
        "channels_with_reports": 0,
        "channels_requested_target_met": 0,
        "channels_capacity_sufficient": 0,
        "channels_partial_below_capacity_goal": 0,
        "channels_with_no_acceptable_probes": 0,
        "channels_with_no_geometric_capacity": 0,
        "channels_missing_reports": len(channels),
        "channels_with_transcript_annotation": 0,
        "channels_transcript_cache_hits": 0,
        "annotation_status": "",
        "cds_start": "",
        "cds_end": "",
        "minimum_final_probe_pairs": "",
        "maximum_final_probe_pairs": "",
        "log": str(log_file),
    }
    return rows, summary


def collect_skipped_outcomes(
    *,
    gene_symbol: str,
    accession: str,
    pipeline_status: str,
    log_file: Path,
    channels: Sequence[str],
) -> tuple[list[dict[str, object]], dict[str, object]]:
    return _collect_rejected_outcomes(
        gene_symbol=gene_symbol,
        accession=accession,
        pipeline_status=pipeline_status,
        report_status="skipped",
        log_file=log_file,
        channels=channels,
    )


def collect_failed_outcomes(
    *,
    gene_symbol: str,
    accession: str,
    pipeline_status: str,
    log_file: Path,
    channels: Sequence[str],
) -> tuple[list[dict[str, object]], dict[str, object]]:
    return _collect_rejected_outcomes(
        gene_symbol=gene_symbol,
        accession=accession,
        pipeline_status=pipeline_status,
        report_status="failed",
        log_file=log_file,
        channels=channels,
    )


def _classify_input_failure(text: str) -> str | None:
    """Classify a deterministic input rejection without hiding bad inputs.

    These messages are emitted only after the organism-scoped NCBI lookup or
    transcript validation has completed. Retryable NCBI/network failures use
    a separate return code and are deliberately not treated as input failures.
    """
    lowered = str(text or "").casefold()
    if "does not belong to the selected organism" in lowered:
        return "skipped_species_mismatch"
    if any(
        marker in lowered
        for marker in (
            "no ncbi gene match was found",
            "no linked refseq rna was found",
            "could not resolve accession",
            "ncbi did not return a valid fasta for",
        )
    ):
        return "failed"
    return None


def _is_species_input_failure(text: str) -> bool:
    """Backward-compatible predicate for true species mismatches only."""
    return _classify_input_failure(text) == "skipped_species_mismatch"


def _manifest_output_roots(
    results_dir: Path,
    gene_dir: str,
    channels: Sequence[str],
    accession: str = "",
    modified_after_ns: int | None = None,
) -> list[str]:
    """Return only target folders belonging to one manifest job.

    A resumed manifest may have no new console file messages.  Searching from
    the manifest root in that case would make the web report show maps from
    unrelated genes, so scope discovery to the current gene and channels.
    """
    roots: list[str] = []
    for channel in channels:
        channel_root = results_dir / gene_dir / f"{gene_dir}_{channel}"
        if not channel_root.is_dir():
            continue
        seen: set[Path] = set()
        for marker in ("*_final_selected_pairs.tsv", "*_final_probe_map.png", "*_final_IDT_order.csv"):
            for path in channel_root.rglob(marker):
                target_root = path.parent
                if accession and target_root.name != f"{gene_dir}_{core.safe_name(accession)}":
                    continue
                if modified_after_ns is not None:
                    try:
                        if path.stat().st_mtime_ns < modified_after_ns:
                            continue
                    except OSError:
                        continue
                if target_root not in seen:
                    roots.append(str(target_root.resolve()))
                    seen.add(target_root)
    return roots


def _manifest_existing_job_matches_species(
    results_dir: Path,
    gene_dir: str,
    channels: Sequence[str],
    species: str,
) -> bool:
    """Check resumable manifest output before trusting its done marker.

    Earlier runs could have cached an accession under a different organism.
    A done marker alone is not enough evidence that the output belongs to the
    selected species, so inspect the stored NCBI transcript titles before
    skipping a job.  A mismatch invalidates only the marker; the normal run
    then refreshes the exact transcript and overwrites the target outputs.
    """
    preset = references.get_species_preset(species)
    expected = preset.scientific_name if preset is not None else ""
    if not expected:
        return True
    found_resolution = False
    for channel in channels:
        channel_root = results_dir / gene_dir / f"{gene_dir}_{channel}"
        for path in channel_root.rglob("*_transcript_resolution.json") if channel_root.is_dir() else ():
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, ValueError, TypeError):
                continue
            if not isinstance(data, dict):
                continue
            found_resolution = True
            candidates = data.get("transcript_candidates") or []
            titles = [str(item.get("title") or "") for item in candidates if isinstance(item, dict)]
            if not titles:
                titles = [str(data.get("title") or "")]
            for title in titles:
                if title and not core._organism_text_matches(title, expected):
                    return False
    return found_resolution


def run_manifest(
    *,
    manifest: Path,
    project_root: Path,
    extra_args: Sequence[str] = (),
    channels: Sequence[str] = CHANNELS,
    smart_default: bool = True,
    progress_callback: ProgressCallback | None = None,
    run_name: str | None = None,
) -> dict[str, object]:
    """Run every valid manifest job inside a dedicated ``manifests`` directory."""
    manifest = Path(manifest).expanduser().resolve()
    project_root = Path(project_root).expanduser().resolve()
    if not manifest.is_file():
        raise ValueError(f"manifest not found: {manifest}")
    selected_channels = tuple(str(channel).strip().upper() for channel in channels)
    if not selected_channels or any(channel not in CHANNELS for channel in selected_channels):
        raise ValueError(f"channels must be drawn from {', '.join(CHANNELS)}")
    preflight_rows, collisions = _preflight_all_rows(manifest)
    jobs = [row for row in preflight_rows if row["status"] in {"valid", "shared_accession"}]
    preflight_invalid = sum(1 for row in preflight_rows if row.get("status") == "invalid")
    species = _option_value(extra_args, "--species", "xtr")
    preset = references.get_species_preset(species)
    if preset is not None:
        species = preset.key
    effective_extra_args = [str(item) for item in extra_args]
    if smart_default and not _option_present(effective_extra_args, "--auto-curate-if-needed"):
        effective_extra_args += ["--auto-curate-if-needed"]
    if smart_default and not _option_present(effective_extra_args, "--target-probes"):
        effective_extra_args += ["--target-probes", "20"]
    resume_payload = _manifest_resume_payload(
        extra_args=effective_extra_args,
        channels=selected_channels,
        smart_default=smart_default,
        species=species,
    )
    resume_signature = _manifest_resume_signature(resume_payload)
    channel_collisions = manifest_channel_collisions(preflight_rows, selected_channels)
    manifest_name = core.safe_name(run_name or manifest.stem) or "manifest"
    manifest_root = references.species_run_root(project_root, species, "manifest") / manifest_name
    status_dir = manifest_root / "status"
    status_tables_dir = status_dir / "tables"
    logs_dir = manifest_root / "logs"
    results_dir = manifest_root / "results"
    status_dir.mkdir(parents=True, exist_ok=True)
    status_tables_dir.mkdir(parents=True, exist_ok=True)
    logs_dir.mkdir(parents=True, exist_ok=True)
    results_dir.mkdir(parents=True, exist_ok=True)
    # Keep an existing resumable catalogue intact when it was created before
    # the status-table folder was introduced.  The named files are moved as a
    # one-time layout migration; resume markers remain directly in status/.
    for table_name in (
        "manifest_preflight.tsv",
        "run_status.tsv",
        "channel_design_outcomes.tsv",
        "transcript_design_summary.tsv",
    ):
        old_path = status_dir / table_name
        new_path = status_tables_dir / table_name
        if old_path.exists() and not new_path.exists():
            old_path.replace(new_path)
    write_preflight(status_tables_dir / "manifest_preflight.tsv", preflight_rows)
    job_channels = [
        (str(row.get("channel") or "").strip().upper(),)
        if str(row.get("channel") or "").strip()
        else selected_channels
        for row in jobs
    ]
    total_units = max(1, sum(len(row_channels) for row_channels in job_channels))
    _notify(
        progress_callback,
        phase="manifest",
        completed=0,
        total=total_units,
        message=f"Manifest preflight complete: {len(jobs)} valid job(s)",
    )
    if collisions or channel_collisions:
        details = "; ".join(
            f"line {row['line_number']} {row['gene_symbol']} -> {row.get('reason', row['gene_dir'])}"
            for row in [*collisions, *channel_collisions]
        )
        raise ValueError(f"manifest output collision detected before design: {details}")

    status_tsv = status_tables_dir / "run_status.tsv"
    _ensure_status_table(status_tsv)
    channel_tsv = status_tables_dir / "channel_design_outcomes.tsv"
    summary_tsv = status_tables_dir / "transcript_design_summary.tsv"
    existing_channel_rows = _read_current_rows(channel_tsv, CHANNEL_OUTCOME_FIELDS)
    existing_summary_rows = _read_current_rows(summary_tsv, SUMMARY_FIELDS)
    channel_updates: dict[tuple[str, ...], list[dict[str, object]]] = {}
    summary_updates: dict[tuple[str, ...], list[dict[str, object]]] = {}
    cache_root = references.species_cache_root(project_root, species, "manifest", manifest_name)
    batch_stdout: list[str] = []
    batch_stderr: list[str] = []
    output_roots: list[str] = []
    completed = 0
    failed = 0
    skipped = 0
    completed_units = 0
    try:
        for job_index, row in enumerate(jobs):
            core.raise_if_cancelled()
            gene_symbol = str(row["gene_symbol"])
            gene_dir = str(row["gene_dir"])
            requested_accession = str(row["accession"])
            row_channels = job_channels[job_index]
            row_units = len(row_channels)
            unit_offset = completed_units
            _notify(
                progress_callback,
                phase="manifest",
                completed=unit_offset,
                total=total_units,
                message=(
                    f"Starting {gene_symbol} ({job_index + 1}/{len(jobs)}); "
                    f"channel(s): {','.join(row_channels)}"
                ),
            )
            explicit_channel = str(row.get("channel") or "").strip().upper() or None
            job_id = _safe_job_name(gene_dir, requested_accession, explicit_channel)
            done_marker = status_dir / f"{job_id}.done"
            failed_marker = status_dir / f"{job_id}.failed"
            skipped_marker = status_dir / f"{job_id}.skipped"
            log_file = logs_dir / f"{job_id}.log"
            if done_marker.exists() and not _manifest_done_marker_matches(done_marker, resume_signature):
                done_marker.unlink(missing_ok=True)
                batch_stdout.append(
                    f"RE-RUN: {gene_symbol} {requested_accession or '(gene-only)'} "
                    "(design options or package version changed)\n"
                )
            if done_marker.exists() and not _manifest_existing_job_matches_species(results_dir, gene_dir, row_channels, species):
                done_marker.unlink(missing_ok=True)
                batch_stdout.append(
                    f"RE-RUN: {gene_symbol} {requested_accession or '(gene-only)'} "
                    "(stored output did not match the selected species)\n"
                )
            if done_marker.exists():
                effective_accession = requested_accession
                rows, summary = collect_design_outcomes(project_root=manifest_root, species=species, gene_symbol=gene_symbol, gene_dir=gene_dir, accession=effective_accession, pipeline_status="complete_existing", log_file=log_file, channels=row_channels)
                output_roots.extend(_manifest_output_roots(results_dir, gene_dir, row_channels, effective_accession))
                key = (gene_symbol, effective_accession, ",".join(row_channels))
                channel_updates[key] = rows
                summary_updates[key] = [summary]
                batch_stdout.append(f"SKIP: {gene_symbol} {requested_accession} ({','.join(row_channels)}; already complete)\n")
                completed_units += row_units
                continue

            failed_marker.unlink(missing_ok=True)
            skipped_marker.unlink(missing_ok=True)
            batch_stdout.append(f"RUN: {gene_symbol} {requested_accession or '(gene-only; automatic accession)'} ({'-'.join(row_channels)})\n")
            job_extra_args = list(extra_args)
            # A manifest may mix gene-only rows with rows that specify an
            # exact accession.  --all-transcripts applies to the former; an
            # explicit accession is already one exact transcript and must not
            # make the whole manifest fail.
            if requested_accession:
                job_extra_args = [item for item in job_extra_args if item != "--all-transcripts"]
            job_started_ns = time.time_ns()
            try:
                result = run_all_channels(
                    gene=gene_symbol,
                    accession=requested_accession or None,
                    output_root=results_dir,
                    extra_args=job_extra_args,
                    channels=row_channels,
                    gene_dir=gene_dir,
                    capture=True,
                    smart_default=smart_default,
                    progress_callback=(
                        lambda update, offset=unit_offset: _notify(
                            progress_callback,
                            phase="manifest",
                            completed=offset + int(update.get("completed", 0)),
                            total=total_units,
                            message=f"{gene_symbol}: {update.get('message', '')}",
                        )
                    ),
                    output_scoped=True,
                    workflow="manifest",
                    cache_root=cache_root,
                )
            except core.RunCancelled:
                raise
            except core.RetryableNCBIError as exc:
                result = {
                    "return_code": 75,
                    "success": False,
                    "resolved_accession": requested_accession,
                    "stdout": "",
                    "stderr": f"ERROR: temporary NCBI failure: {exc}\n",
                }
            except (OSError, RuntimeError, ValueError) as exc:
                # Keep one malformed or failed row from aborting the rest of
                # the catalogue.  The row remains visible in status/logs.
                result = {
                    "return_code": 1,
                    "success": False,
                    "resolved_accession": requested_accession,
                    "stdout": "",
                    "stderr": f"ERROR: {exc}\n",
                }
            effective_accession = str(result.get("resolved_accession") or requested_accession)
            text = str(result.get("stdout", "")) + str(result.get("stderr", ""))
            log_file.write_text(text, encoding="utf-8")
            batch_stdout.append(str(result.get("stdout", "")))
            batch_stderr.append(str(result.get("stderr", "")))
            if bool(result.get("success")):
                done_marker.write_text(
                    json.dumps(
                        {
                            "signature": resume_signature,
                            "parameters": resume_payload,
                            "completed_at": utc_timestamp(),
                        },
                        indent=2,
                    )
                    + "\n",
                    encoding="utf-8",
                )
                failed_marker.unlink(missing_ok=True)
                skipped_marker.unlink(missing_ok=True)
                status = "complete"
                completed += 1
            elif int(result.get("return_code", 1)) != 75 and _classify_input_failure(text):
                # The core validates NCBI transcript ownership and resolution
                # before writing transcript FASTA/cache data.  A true species
                # mismatch is a non-fatal skip; an unresolved gene/accession
                # is a failed row so a typo cannot pass silently.
                input_status = _classify_input_failure(text)
                done_marker.unlink(missing_ok=True)
                failed_marker.unlink(missing_ok=True)
                if input_status == "skipped_species_mismatch":
                    skipped_marker.write_text(text or "species-specific input was rejected\n", encoding="utf-8")
                    status = input_status
                    skipped += 1
                    batch_stderr.append(
                        f"SKIPPED: {gene_symbol} {requested_accession or '(gene-only)'} "
                        f"is not available for {species}; no design or cache output was kept.\n"
                    )
                else:
                    failed_marker.write_text(text or "input could not be resolved\n", encoding="utf-8")
                    status = "failed"
                    failed += 1
                    batch_stderr.append(
                        f"FAILED INPUT: {gene_symbol} {requested_accession or '(gene-only)'} "
                        "could not be resolved; correct the manifest row and rerun.\n"
                    )
            else:
                done_marker.unlink(missing_ok=True)
                skipped_marker.unlink(missing_ok=True)
                failed_marker.write_text(str(result.get("return_code", 1)), encoding="utf-8")
                status = "failed"
                failed += 1
            with status_tsv.open("a", encoding="utf-8") as handle:
                handle.write(
                    f"{utc_timestamp()}\t{gene_symbol}\t{effective_accession}\t"
                    f"{','.join(row_channels)}\t{status}\t{log_file}\n"
                )
            if status == "skipped_species_mismatch":
                rows, summary = collect_skipped_outcomes(
                    gene_symbol=gene_symbol,
                    accession=effective_accession,
                    pipeline_status=status,
                    log_file=log_file,
                    channels=row_channels,
                )
            elif status == "failed" and not bool(result.get("success")):
                rows, summary = collect_failed_outcomes(
                    gene_symbol=gene_symbol,
                    accession=effective_accession,
                    pipeline_status=status,
                    log_file=log_file,
                    channels=row_channels,
                )
            else:
                rows, summary = collect_design_outcomes(project_root=manifest_root, species=species, gene_symbol=gene_symbol, gene_dir=gene_dir, accession=effective_accession, pipeline_status=status, log_file=log_file, channels=row_channels)
                output_roots.extend(
                    _manifest_output_roots(
                        results_dir,
                        gene_dir,
                        row_channels,
                        effective_accession,
                        modified_after_ns=job_started_ns,
                    )
                )
            key = (gene_symbol, effective_accession, ",".join(row_channels))
            channel_updates[key] = rows
            summary_updates[key] = [summary]
            if status == "failed":
                batch_stderr.append(f"FAILED: {gene_symbol} {effective_accession}; see {log_file}\n")
            completed_units += row_units
    finally:
        _write_current_row_updates(channel_tsv, CHANNEL_OUTCOME_FIELDS, existing_channel_rows, channel_updates)
        _write_current_row_updates(summary_tsv, SUMMARY_FIELDS, existing_summary_rows, summary_updates)
    return {
        "return_code": 0 if failed == 0 and preflight_invalid == 0 else 1,
        "success": failed == 0 and preflight_invalid == 0,
        "jobs": len(jobs),
        "preflight_skipped": len(preflight_rows) - len(jobs),
        "preflight_invalid": preflight_invalid,
        "completed": completed,
        "failed": failed,
        "skipped": skipped,
        "stdout": "".join(batch_stdout),
        "stderr": "".join(batch_stderr),
        "project_root": str(manifest_root),
        "roots": list(dict.fromkeys(output_roots)),
    }


def _channels(value: str) -> tuple[str, ...]:
    chosen = tuple(item.strip().upper() for item in value.split(",") if item.strip())
    if not chosen or any(item not in CHANNELS for item in chosen):
        raise argparse.ArgumentTypeError(f"channels must be a comma-separated subset of {','.join(CHANNELS)}")
    return chosen


def all_channels_main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run HCRProbeForge for one target across selected HCR channels.")
    parser.add_argument("gene", nargs="?", help="Gene label; may contain filesystem-unsafe characters")
    parser.add_argument("accession_pos", nargs="?", help="Optional RefSeq accession")
    parser.add_argument("--accession", dest="accession_opt", help="Optional RefSeq accession")
    parser.add_argument("--fasta", type=Path, help="Single-record FASTA input")
    parser.add_argument("--outdir", type=Path, default=Path("hcr_results"), help="Project output root; runs and cache are created below it")
    parser.add_argument("--channels", type=_channels, default=CHANNELS, help="Comma-separated channels; default: B1,B2,B3,B4,B5")
    parser.add_argument("--version", action="version", version=f"%(prog)s {core.__version__}")
    args, extra = parser.parse_known_args(list(argv) if argv is not None else None)
    accession = args.accession_opt or args.accession_pos
    if args.accession_opt and args.accession_pos:
        parser.error("provide accession either positionally or with --accession, not both")
    try:
        result = run_all_channels(gene=args.gene, accession=accession, fasta=args.fasta, output_root=args.outdir, extra_args=extra, channels=args.channels, capture=False, smart_default=True)
    except ValueError as exc:
        parser.error(str(exc))
    return int(result["return_code"])


def manifest_main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run HCRProbeForge for every gene/accession/channel row in a manifest TSV/CSV/text file.")
    parser.add_argument("manifest", type=Path, help="Manifest with gene, optional accession, and optional B1-B5 channel columns")
    parser.add_argument("project_root", type=Path, help="Catalogue/project output directory")
    parser.add_argument("--channels", type=_channels, default=CHANNELS)
    parser.add_argument("--version", action="version", version=f"%(prog)s {core.__version__}")
    args, extra = parser.parse_known_args(list(argv) if argv is not None else None)
    try:
        result = run_manifest(manifest=args.manifest, project_root=args.project_root, extra_args=extra, channels=args.channels)
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(str(result.get("stdout", "")), end="")
    if result.get("stderr"):
        print(str(result["stderr"]), end="", file=sys.stderr)
    print(
        f"Batch finished: jobs={result['jobs']} completed={result['completed']} "
        f"skipped={result.get('skipped', 0)} preflight_skipped={result.get('preflight_skipped', 0)} "
        f"preflight_invalid={result.get('preflight_invalid', 0)} "
        f"failed={result['failed']}"
    )
    return int(result["return_code"])


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "manifest":
        raise SystemExit(manifest_main(sys.argv[2:]))
    if len(sys.argv) > 1 and sys.argv[1] == "all-channels":
        raise SystemExit(all_channels_main(sys.argv[2:]))
    raise SystemExit(all_channels_main())
